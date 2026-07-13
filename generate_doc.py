from docx import Document
from docx.shared import Pt, Inches

doc = Document()
doc.add_heading('End-to-End Sales Forecasting & Demand Intelligence System', 0)

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    "This report presents an end-to-end sales forecasting and demand intelligence system designed to optimize stocking and inventory decisions. "
    "By analyzing historical sales data, detecting demand anomalies, and predicting future trends across multiple categories and regions, "
    "we are equipped to minimize both overstock and stockout scenarios. The selected SARIMA model provides robust monthly forecasts, enabling "
    "proactive rather than reactive supply chain management."
)

doc.add_heading('2. Key Findings from EDA and Forecasting', level=1)
doc.add_paragraph(
    "- The Technology category consistently generates the highest total revenue among all products.\n"
    "- The West region exhibits the strongest and most consistent sales growth over the four-year period.\n"
    "- Sales display strong seasonality, with consistent demand spikes occurring every November and December, likely due to holiday promotions.\n"
    "- The SARIMA model was identified as the best-performing forecast model for monthly sales due to its low MAPE compared to Prophet and XGBoost."
)

doc.add_heading('3. 3-Month Sales Forecast (Plain Language)', level=1)
doc.add_paragraph(
    "Based on the best model, overall sales for the upcoming three months are projected to show stability with seasonal trends intact. "
    "Forecast values predict a steady baseline demand across major product categories. We are 95% confident that actual sales will fall "
    "within our predicted ranges, allowing the supply chain team to plan stock replenishment with minimal risk."
)

doc.add_heading('4. Top 3 Anomalies Detected and Likely Causes', level=1)
doc.add_paragraph(
    "1. Late November Spikes: Extreme sales spikes consistently flagged in late November, which strongly correlates with Black Friday and Thanksgiving holiday sales.\n"
    "2. Mid-December Spikes: High volumes just before Christmas, capturing last-minute holiday shopping.\n"
    "3. Low-Volume January Drops: A significant dip in sales immediately following the holiday season, representing the natural post-holiday demand exhaustion."
)

doc.add_heading('5. Product Demand Segmentation', level=1)
doc.add_paragraph(
    "Using K-Means clustering, products were grouped into 4 distinct segments based on total sales volume, average order value, and volatility:\n"
    "- High Volume, Stable Demand: Focus on automated replenishment.\n"
    "- Low Volume, Low Volatility: Requires minimal buffer stock; consider dropshipping.\n"
    "- Moderate Volume, Moderate Volatility: Needs close tracking of seasonal fluctuations.\n"
    "- High Volatility, High Demand: Highly profitable but risky; needs large safety buffers during peak periods."
)

doc.add_heading('6. Business Recommendations', level=1)
doc.add_paragraph(
    "1. Increase holiday safety stock for Technology products by 20% starting in October to capture consistent November/December surges.\n"
    "2. Expand West region warehouse capacity to support its high and growing volume, reducing reliance on long-distance shipping.\n"
    "3. Implement automated reordering for 'High Volume, Stable Demand' products, freeing up buyer resources to actively manage volatile categories."
)

doc.add_heading('7. Risks and Limitations', level=1)
doc.add_paragraph(
    "One major limitation of this statistical forecasting system is its inability to account for unprecedented macroeconomic shocks "
    "(such as pandemics or sudden inflation) without external covariates. Future iterations should incorporate external market indicators "
    "to improve resilience against sudden structural shifts."
)

doc.save('summary.docx')
